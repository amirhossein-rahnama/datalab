#!/usr/bin/env python3
"""Markdown outline -> PPTX via python-pptx.

Outline format:
  '# '           begins a new content slide (its title).
  '- ' / '* '    or plain lines are bullets on the current content slide.
  '![alt](path)' becomes a figure slide (image scaled to fit; path resolved
                 relative to the outline file).
  '| a | b |'    a GitHub-style table block becomes a table slide.
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

IMG_RE = re.compile(r"^!\[[^\]]*\]\(([^)]+)\)\s*$")


def _parse_table(lines):
    rows = []
    for ln in lines:
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        if cells and all(c and set(c) <= set("-: ") for c in cells):
            continue  # separator row like |---|:--:|
        rows.append(cells)
    return rows


def _parse(outline_text: str):
    """Parse the outline into an ordered list of blocks:
    {"kind": "content", "title", "bullets"} | {"kind": "image", "title", "path"}
    | {"kind": "table", "title", "rows"}."""
    blocks = []
    current_title = ""
    content = None
    table_lines: list[str] = []

    def flush_table():
        nonlocal table_lines
        if table_lines:
            rows = _parse_table(table_lines)
            if rows:
                blocks.append({"kind": "table", "title": current_title, "rows": rows})
            table_lines = []

    def flush_content():
        nonlocal content
        if content is not None:
            if content["bullets"]:  # skip redundant title-only slides (section is just a table/figure)
                blocks.append(content)
            content = None

    for raw in outline_text.splitlines():
        line = raw.rstrip()
        s = line.strip()
        if not s:
            continue
        if line.startswith("# "):
            flush_table(); flush_content()
            current_title = line[2:].strip()
            content = {"kind": "content", "title": current_title, "bullets": []}
        elif IMG_RE.match(s):
            flush_table(); flush_content()
            blocks.append({"kind": "image", "title": current_title, "path": IMG_RE.match(s).group(1)})
        elif s.startswith("|") and s.endswith("|"):
            flush_content()
            table_lines.append(s)
        elif s.startswith(("- ", "* ")):
            flush_table()
            if content is None:
                content = {"kind": "content", "title": current_title, "bullets": []}
            content["bullets"].append(s[2:])
        else:
            flush_table()
            if content is None:
                current_title = s
                content = {"kind": "content", "title": s, "bullets": []}
            else:
                content["bullets"].append(s)
    flush_table(); flush_content()
    return blocks


def _add_content_slide(prs, block):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = block["title"]
    body = slide.placeholders[1].text_frame
    body.clear()
    for j, bullet in enumerate(block["bullets"]):
        para = body.paragraphs[0] if j == 0 else body.add_paragraph()
        para.text = bullet
        para.font.size = Pt(18)


def _add_image_slide(prs, block, base_dir):
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    slide.shapes.title.text = block["title"] or "Figure"
    path = Path(block["path"])
    if not path.is_absolute():
        path = Path(base_dir) / block["path"]
    if not path.exists():
        return  # missing image: keep the titled slide, skip the picture
    pic = slide.shapes.add_picture(str(path), 0, 0, width=Inches(8))
    max_h = Inches(5.3)
    if pic.height > max_h:
        ratio = max_h / pic.height
        pic.height = int(pic.height * ratio)
        pic.width = int(pic.width * ratio)
    pic.left = int((prs.slide_width - pic.width) / 2)
    pic.top = Inches(1.7)


def _add_table_slide(prs, block):
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    slide.shapes.title.text = block["title"] or "Table"
    rows = block["rows"]
    nrows, ncols = len(rows), max(len(r) for r in rows)
    table = slide.shapes.add_table(
        nrows, ncols, Inches(0.7), Inches(1.7), Inches(8.6), Inches(0.4 * nrows)
    ).table
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.text = row[j] if j < len(row) else ""
            for para in cell.text_frame.paragraphs:
                para.font.size = Pt(16 if i == 0 else 14)


def outline_to_pptx(outline_text: str, output_path, title: str | None = None, base_dir=None) -> Path:
    output_path = Path(output_path)
    base_dir = Path(base_dir) if base_dir else Path.cwd()
    prs = Presentation()
    if title:
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = title
    for block in _parse(outline_text):
        if block["kind"] == "content":
            _add_content_slide(prs, block)
        elif block["kind"] == "image":
            _add_image_slide(prs, block, base_dir)
        elif block["kind"] == "table":
            _add_table_slide(prs, block)
    prs.save(str(output_path))
    return output_path


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Markdown outline -> PPTX")
    ap.add_argument("--input", required=True)
    ap.add_argument("--output", required=True)
    ap.add_argument("--title", default=None)
    args = ap.parse_args(argv)
    md_path = Path(args.input)
    text = md_path.read_text(encoding="utf-8")
    out = outline_to_pptx(text, args.output, args.title, base_dir=md_path.parent)
    print(f"Wrote {out} ({out.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
