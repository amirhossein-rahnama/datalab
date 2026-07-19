import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import generate_deck  # noqa: E402

OUTLINE = "# Problem\n- context\n- goal\n\n# Results\n- accuracy 0.92\n- beats baseline\n"

def test_parse_splits_slides():
    slides = generate_deck._parse(OUTLINE)
    assert [s["title"] for s in slides] == ["Problem", "Results"]
    assert slides[0]["bullets"] == ["context", "goal"]

def test_produces_valid_pptx(tmp_path):
    from pptx import Presentation
    out = tmp_path / "deck.pptx"
    result = generate_deck.outline_to_pptx(OUTLINE, out, title="My Deck")
    assert result == out and out.exists() and out.stat().st_size > 0
    prs = Presentation(str(out))
    # title slide + 2 content slides
    assert len(prs.slides) == 3


def test_deck_embeds_table_and_image(tmp_path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    img = tmp_path / "fig.png"
    fig = plt.figure()
    plt.plot([0, 1], [1, 0])
    fig.savefig(img)
    plt.close(fig)

    outline = "# Results\n| Metric | Value |\n|---|---|\n| Acc | 0.9 |\n\n![chart](fig.png)\n"
    out = tmp_path / "deck.pptx"
    generate_deck.outline_to_pptx(outline, out, title="T", base_dir=tmp_path)

    shapes = [sh for slide in Presentation(str(out)).slides for sh in slide.shapes]
    assert any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sh in shapes)
    assert any(getattr(sh, "has_table", False) for sh in shapes)


def test_deck_missing_image_does_not_crash(tmp_path):
    from pptx import Presentation
    out = tmp_path / "deck.pptx"
    generate_deck.outline_to_pptx("# S\n![x](nope.png)\n", out, base_dir=tmp_path)
    assert list(Presentation(str(out)).slides)  # produced a deck, no crash
